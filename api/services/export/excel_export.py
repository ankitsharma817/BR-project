"""
Generate an Excel (.xlsx) report for match results.
Uses openpyxl — install: pip install openpyxl
"""
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, numbers,
)
from openpyxl.utils import get_column_letter

from ...models.matching import MatchingResult, MatchAnalysis
from ...models.proposal import Proposal

BLUE = "1677FF"
GREEN = "52C41A"
YELLOW = "FAAD14"
ORANGE = "FA8C16"
RED = "F5222D"
LIGHT_BLUE = "E6F4FF"
GREY = "F0F2F5"


def _score_hex(score: float) -> str:
    if score >= 0.80: return GREEN
    if score >= 0.60: return YELLOW
    if score >= 0.30: return ORANGE
    return RED


def _header_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _thin_border() -> Border:
    side = Side(style="thin", color="D9D9D9")
    return Border(left=side, right=side, top=side, bottom=side)


def _set_col_widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def generate_match_excel(
    proposal: Proposal,
    result: MatchingResult,
    analysis: MatchAnalysis | None,
) -> bytes:
    wb = Workbook()

    # ── Sheet 1: Summary ─────────────────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Summary"
    _set_col_widths(ws_sum, [28, 18, 18])

    ws_sum.append(["BR Matching Report"])
    ws_sum["A1"].font = Font(bold=True, size=16, color=BLUE)
    ws_sum.append([f"Vendor: {proposal.vendor_name}"])
    ws_sum.append([f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"])
    ws_sum.append([f"Report Version: {result.version}"])
    ws_sum.append([])

    # Overall score row
    ws_sum.append(["OVERALL SCORE", f"{round(result.overall_score * 100, 1)}%"])
    ws_sum["A6"].font = Font(bold=True, size=13)
    ws_sum["B6"].font = Font(bold=True, size=13, color=_score_hex(result.overall_score))

    ws_sum.append([])
    ws_sum.append(["Category", "Score", "Status"])
    for cell in ws_sum[8]:
        cell.fill = _header_fill(BLUE)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = _thin_border()

    cats = [
        ("Functional", result.functional_score),
        ("Technical", result.technical_score),
        ("Compliance", result.compliance_score),
        ("Security", result.security_score),
        ("Timeline", result.timeline_score),
        ("Resource", result.resource_score),
        ("Deliverables", result.deliverables_score),
    ]
    for row_i, (name, score) in enumerate(cats, 9):
        if score is None:
            continue
        pct = f"{round(score * 100, 1)}%"
        status = "Strong" if score >= 0.8 else "Partial" if score >= 0.6 else "Gap" if score >= 0.3 else "Missing"
        ws_sum.append([name, pct, status])
        score_cell = ws_sum.cell(row=row_i, column=2)
        score_cell.font = Font(color=_score_hex(score), bold=True)
        for col in range(1, 4):
            ws_sum.cell(row=row_i, column=col).border = _thin_border()

    if result.executive_summary:
        ws_sum.append([])
        ws_sum.append(["Executive Summary"])
        ws_sum[ws_sum.max_row]["A"].font = Font(bold=True)
        ws_sum.append([result.executive_summary])
        ws_sum.cell(ws_sum.max_row, 1).alignment = Alignment(wrap_text=True)
        ws_sum.row_dimensions[ws_sum.max_row].height = 60

    # ── Sheet 2: Requirement Breakdown ───────────────────────────────────────────
    ws_req = wb.create_sheet("Requirements")
    _set_col_widths(ws_req, [55, 15, 10, 10, 18, 60])
    headers = ["Requirement", "Category", "Priority", "Score %", "Status", "Explanation"]
    ws_req.append(headers)
    for i, cell in enumerate(ws_req[1], 1):
        cell.fill = _header_fill(BLUE)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = _thin_border()

    for row_i, rm in enumerate(result.requirement_matchings, 2):
        br = rm.br_requirement
        prop = rm.proposal_requirement
        ws_req.append([
            br.text if br else "",
            br.category if br else "",
            br.priority if br else "",
            round(rm.score * 100, 1),
            rm.label.replace("_", " ").title(),
            rm.explanation or "",
        ])
        score_cell = ws_req.cell(row=row_i, column=4)
        score_cell.font = Font(color=_score_hex(rm.score), bold=True)
        fill_color = LIGHT_BLUE if row_i % 2 == 0 else "FFFFFF"
        for col in range(1, 7):
            c = ws_req.cell(row=row_i, column=col)
            c.fill = PatternFill("solid", fgColor=fill_color)
            c.border = _thin_border()
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws_req.row_dimensions[row_i].height = 40

    # ── Sheet 3: Analysis ────────────────────────────────────────────────────────
    if analysis:
        ws_ana = wb.create_sheet("Analysis")
        _set_col_widths(ws_ana, [15, 60, 20])

        def _section(title: str, rows: list[list]):
            ws_ana.append([title])
            ws_ana[ws_ana.max_row][0].font = Font(bold=True, size=12, color=BLUE)
            if rows:
                ws_ana.append(rows[0])  # header row
                for cell in ws_ana[ws_ana.max_row]:
                    cell.fill = _header_fill(BLUE)
                    cell.font = Font(bold=True, color="FFFFFF")
                for r in rows[1:]:
                    ws_ana.append(r)
            ws_ana.append([])

        risk_rows = [["Level", "Description", "Category"]]
        for r in (analysis.risks or []):
            risk_rows.append([
                r.get("level", "").upper() if isinstance(r, dict) else r.level.upper(),
                r.get("description", "") if isinstance(r, dict) else r.description,
                r.get("category", "") if isinstance(r, dict) else (r.category or ""),
            ])
        _section("Risks", risk_rows)

        rec_rows = [["Priority", "Action", "Reason"]]
        for rec in (analysis.recommendations or []):
            rec_rows.append([
                rec.get("priority", "") if isinstance(rec, dict) else rec.priority,
                rec.get("action", "") if isinstance(rec, dict) else rec.action,
                rec.get("reason", "") if isinstance(rec, dict) else (rec.reason or ""),
            ])
        _section("Recommendations", rec_rows)

        ws_ana.append(["Strengths"])
        ws_ana[ws_ana.max_row][0].font = Font(bold=True, size=12, color=BLUE)
        for s in (analysis.strengths or []):
            ws_ana.append(["", s])

        ws_ana.append([])
        ws_ana.append(["Gaps"])
        ws_ana[ws_ana.max_row][0].font = Font(bold=True, size=12, color=RED)
        for g in (analysis.gaps or []):
            ws_ana.append(["", g])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_comparison_excel(comparisons: list[dict]) -> bytes:
    """Side-by-side comparison of multiple proposals."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    cats = ["functional", "technical", "compliance", "security", "timeline", "resource", "deliverables"]
    headers = ["Category"] + [c["vendor_name"] for c in comparisons]
    _set_col_widths(ws, [20] + [22] * len(comparisons))
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _header_fill(BLUE)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = _thin_border()

    # Overall row
    overall_row = ["OVERALL"] + [
        f"{round((c['overall_score'] or 0) * 100, 1)}%" if c.get("overall_score") is not None else "—"
        for c in comparisons
    ]
    ws.append(overall_row)
    ws[ws.max_row][0].font = Font(bold=True)
    for i, c in enumerate(comparisons, 2):
        score = c.get("overall_score")
        if score is not None:
            cell = ws.cell(ws.max_row, i)
            cell.font = Font(bold=True, color=_score_hex(score))

    for cat in cats:
        row = [cat.capitalize()]
        for c in comparisons:
            score = c.get("category_scores", {}).get(cat)
            row.append(f"{round(score * 100, 1)}%" if score is not None else "—")
        ws.append(row)
        for i, c in enumerate(comparisons, 2):
            score = c.get("category_scores", {}).get(cat)
            if score is not None:
                cell = ws.cell(ws.max_row, i)
                cell.font = Font(color=_score_hex(score))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
